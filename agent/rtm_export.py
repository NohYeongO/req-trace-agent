"""검수용 요구사항 추적표 엑셀 — 검수·대금 청구 때 제출하는 문서.

시트 3개:
1. 표지·요약     사업명, 기준일, 진척률, 상태·분류별 건수, 미완료 목록, 검수 확인란
2. 요구사항 추적표  요구사항별 완료 조건(충족 표시), 상태, 완료일, 담당자, 구현 근거와 확인 결과, 관련 보고
3. 완료 보고 이력  기준일까지 개발자가 확인한 보고

사람이 확인한 데이터만 넣는다 — 확정된 완료 조건과 개발자가 확인한 보고만 쓰고, AI 초안
단계의 조건은 '조건 미확정'으로만 표시한다. 숫자는 대시보드와 같은 build_dashboard에서 온다.
"""
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .dashboard import build_dashboard
from .store.db import DONE, NOT_STARTED, PARTIAL, STATUSES, UNDEFINED, ProjectStore

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
STATUS_FILL = {DONE: "E2F0D9", PARTIAL: "FFF2CC", NOT_STARTED: "F8CBAD", UNDEFINED: "D9D9D9"}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
EV_LABEL = {"api": "API", "file": "파일", "code": "코드", "test": "테스트", "pr": "PR", "commit": "커밋"}

TRACE_COLUMNS = [  # (헤더, 너비)
    ("요구사항ID", 11), ("분류", 7), ("요구사항명", 20), ("상세설명", 36), ("우선순위", 7),
    ("완료 조건", 44), ("충족/전체", 9), ("상태", 9), ("완료일", 11), ("담당자", 12),
    ("구현 근거", 44), ("근거 확인", 12), ("관련 보고", 10), ("검수 확인", 10), ("비고", 16),
]


def _header(ws, row: int, values: list[str]) -> None:
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _write_row(ws, row: int, values: list, fill: str | None = None) -> None:
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.alignment = WRAP
        cell.border = BORDER
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)


def _print_setup(ws, title_rows: str | None = None) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if title_rows:
        ws.print_title_rows = title_rows


def _summary_sheet(ws, dash: dict, project_name: str) -> None:
    ws.title = "표지·요약"
    o = dash["overall"]
    ws["A1"] = "요구사항 추적표"
    ws["A1"].font = Font(bold=True, size=18)
    info = [("사업명", project_name), ("기준일", dash["as_of"]),
            ("작성일", datetime.now().strftime("%Y-%m-%d %H:%M")),
            ("요구사항 수", o["requirements"]),
            ("진척률(요구사항 기준)", f"{o['requirement_rate']}%  (완료 {o['counts'][DONE]} / {o['defined']}건)"),
            ("진척률(완료 조건 기준)", f"{o['criteria_rate']}%  (충족 {o['criteria_met']} / {o['criteria_total']}개)"),
            ("반영된 완료 보고", f"{o['reports']}건")]
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    r = len(info) + 4

    ws.cell(row=r, column=1, value="상태별 현황").font = Font(bold=True, size=12)
    _header(ws, r + 1, ["상태", "건수"])
    for i, s in enumerate(STATUSES, start=r + 2):
        _write_row(ws, i, [s, o["counts"][s]], STATUS_FILL[s])
    r += len(STATUSES) + 3

    ws.cell(row=r, column=1, value="분류별 현황").font = Font(bold=True, size=12)
    _header(ws, r + 1, ["분류", "전체", DONE, PARTIAL, NOT_STARTED, UNDEFINED, "조건 충족률"])
    for i, g in enumerate(dash["by_category"], start=r + 2):
        _write_row(ws, i, [g["category"], g["total"], g[DONE], g[PARTIAL], g[NOT_STARTED],
                           g[UNDEFINED], f"{g['criteria_rate']}%"])
    r += len(dash["by_category"]) + 3

    pending = [x for x in dash["requirements"] if x["status"] != DONE]
    ws.cell(row=r, column=1, value=f"미완료 요구사항 ({len(pending)}건)").font = Font(bold=True, size=12)
    _header(ws, r + 1, ["요구사항ID", "요구사항명", "상태", "충족/전체", "우선순위"])
    for i, x in enumerate(pending, start=r + 2):
        _write_row(ws, i, [x["id"], x["name"], x["status"], f"{x['met']}/{x['total']}", x["priority"]],
                   STATUS_FILL[x["status"]])
    r += len(pending) + 3

    ws.cell(row=r, column=1, value="검수 확인").font = Font(bold=True, size=12)
    _header(ws, r + 1, ["구분", "소속", "성명", "확인일", "서명"])
    for i, role in enumerate(["발주기관 검수자", "수행사 PM"], start=r + 2):
        _write_row(ws, i, [role, "", "", "", ""])
        ws.row_dimensions[i].height = 28
    for col, width in zip("ABCDEFG", (22, 40, 10, 10, 10, 10, 12)):
        ws.column_dimensions[col].width = width
    _print_setup(ws)


def _evidence_text(items: list[dict]) -> tuple[str, str]:
    """(근거 목록 문자열, 확인 결과 요약). 같은 근거가 여러 보고에 나오면 한 번만."""
    seen, lines, found, checkable = set(), [], 0, 0
    for e in items:
        key = (e["type"], e["value"])
        if key in seen:
            continue
        seen.add(key)
        mark = {"found": "✔", "not_found": "✘"}.get(e.get("status"), "–")
        lines.append(f"{mark} [{EV_LABEL.get(e['type'], e['type'])}] {e['value']}")
        if e.get("status") in ("found", "not_found"):
            checkable += 1
            found += e["status"] == "found"
    return "\n".join(lines), (f"{found}/{checkable} 확인" if checkable else "확인 대상 없음")


def _trace_sheet(ws, store: ProjectStore, dash: dict, reports: list[dict]) -> None:
    ws.title = "요구사항 추적표"
    _header(ws, 1, [h for h, _ in TRACE_COLUMNS])
    for i, (_, w) in enumerate(TRACE_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    authors, evidence, report_ids, met_ids = defaultdict(list), defaultdict(list), defaultdict(list), defaultdict(set)
    for rep in reports:
        for link in rep["links"]:
            rid = link["req_id"]
            if rep["author"] not in authors[rid]:
                authors[rid].append(rep["author"])
            evidence[rid] += link["evidence_check"] or [{**e, "status": None} for e in link["evidence"]]
            report_ids[rid].append(f"#{rep['id']}")
            met_ids[rid] |= {c["id"] for c in link["met_criteria"]}

    prog = {x["id"]: x for x in dash["requirements"]}
    for row, req in enumerate((store.get_requirement(x["id"]) for x in dash["requirements"]), start=2):
        p = prog[req["id"]]
        confirmed = [c for c in req["criteria"] if c["confirmed"]]
        crit_text = "\n".join(f"{'✔' if c['id'] in met_ids[req['id']] else '☐'} {c['text']}"
                              for c in confirmed) or "(완료 조건 미확정)"
        ev_text, ev_check = _evidence_text(evidence[req["id"]])
        _write_row(ws, row, [
            req["id"], req["category"], req["name"], req["description"], req["priority"],
            crit_text, f"{p['met']}/{p['total']}", p["status"],
            (p["completed_at"] or "")[:10], ", ".join(authors[req["id"]]),
            ev_text, ev_check if evidence[req["id"]] else "", ", ".join(report_ids[req["id"]]), "", ""])
        ws.cell(row=row, column=8).fill = PatternFill("solid", fgColor=STATUS_FILL[p["status"]])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(TRACE_COLUMNS))}{len(dash['requirements']) + 1}"
    _print_setup(ws, "1:1")


def _history_sheet(ws, reports: list[dict]) -> None:
    ws.title = "완료 보고 이력"
    cols = [("보고ID", 8), ("일시", 17), ("작성자", 10), ("요약", 40), ("연결 요구사항", 14),
            ("충족 조건", 50), ("보고 원문", 60)]
    _header(ws, 1, [h for h, _ in cols])
    for i, (_, w) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row, rep in enumerate(reports, start=2):
        _write_row(ws, row, [
            rep["id"], rep["created_at"].replace("T", " ")[:16], rep["author"], rep["summary"],
            "\n".join(l["req_id"] for l in rep["links"]),
            "\n".join(f"{l['req_id']}: {c['text']}" for l in rep["links"] for c in l["met_criteria"]),
            rep["raw_text"]])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{len(reports) + 1}"
    _print_setup(ws, "1:1")


def export_rtm(store: ProjectStore, out_path: Path, project_name: str,
               as_of: str | None = None, today: date | None = None) -> dict:
    """추적표 엑셀을 쓰고, 엑셀에 쓴 숫자(대시보드와 같은 값)를 돌려준다."""
    dash = build_dashboard(store, as_of, today=today)
    reports = sorted(store.list_reports(limit=100000, until=as_of), key=lambda r: (r["created_at"], r["id"]))
    wb = Workbook()
    _summary_sheet(wb.active, dash, project_name)
    _trace_sheet(wb.create_sheet(), store, dash, reports)
    _history_sheet(wb.create_sheet(), reports)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {"path": str(out_path), "as_of": dash["as_of"], "overall": dash["overall"]}
