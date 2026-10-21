"""산출물 문서를 xlsx로 쓴다 — 인터페이스·테이블 정의서처럼 표가 많은 문서는 엑셀로 많이 제출한다.

표 섹션마다 시트 하나, 같은 group(인터페이스 상세·테이블 정의)은 한 시트에 제목 줄을 두고 이어 쓴다.
서술 섹션은 '종합' 시트에 모은다.
"""
import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_HEAD = PatternFill("solid", fgColor="DCE6F1")
_TITLE = Font(bold=True, size=12)
_THIN = Side(style="thin", color="999999")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
GROUP_NAMES = {"interface_detail": "인터페이스 상세", "table_detail": "테이블 정의"}


def _sheet_name(title: str, used: set) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "", re.sub(r"^\d+(-\d+)?\.\s*", "", title))[:28] or "시트"
    base, i = name, 2
    while name in used:
        name, i = f"{base[:25]}_{i}", i + 1
    used.add(name)
    return name


def _write_table(ws, row: int, sec: dict) -> int:
    for j, c in enumerate(sec["columns"], 1):
        cell = ws.cell(row, j, c)
        cell.font, cell.fill, cell.border = Font(bold=True), _HEAD, _BOX
    for r in sec["rows"] or [["해당 없음"]]:
        row += 1
        for j, v in enumerate(r, 1):
            cell = ws.cell(row, j, v)
            cell.border, cell.alignment = _BOX, Alignment(wrap_text=True, vertical="top")
    if sec.get("note"):
        row += 1
        ws.cell(row, 1, f"※ {sec['note']}").font = Font(size=9, color="808080")
    return row + 2


def render_xlsx(record: dict) -> bytes:
    content = record["content"]
    wb = Workbook()
    cover = wb.active
    cover.title = "표지"
    cover["A1"], cover["A1"].font = content["title"], Font(bold=True, size=16)
    cover["A2"] = content["meta"].get("project_name", "")
    cover["A3"] = f"상태: {'확정' if record['status'] == 'confirmed' else '초안(검토 전)'}" + (
        f" · 확정 {record.get('confirmed_by')} {(record.get('confirmed_at') or '')[:16]}" if record.get("confirmed_by") else "")
    cover["A4"] = f"작성일: {content['meta'].get('written_on', '')}"
    used, sheets = {"표지"}, {}
    texts = []
    for sec in content["sections"]:
        if sec["kind"] != "table":
            texts.append(sec)
            continue
        group = sec.get("group")
        if group:
            if group not in sheets:
                sheets[group] = [wb.create_sheet(_sheet_name(GROUP_NAMES.get(group, group), used)), 1]
            ws, row = sheets[group]
            ws.cell(row, 1, sec["title"]).font = _TITLE
            sheets[group][1] = _write_table(ws, row + 1, sec)
        else:
            ws = wb.create_sheet(_sheet_name(sec["title"], used))
            ws.cell(1, 1, sec["title"]).font = _TITLE
            _write_table(ws, 2, sec)
    if texts:
        ws = wb.create_sheet(_sheet_name("종합", used))
        row = 1
        for sec in texts:
            ws.cell(row, 1, sec["title"]).font = _TITLE
            ws.cell(row + 1, 1, sec.get("text") or "").alignment = Alignment(wrap_text=True, vertical="top")
            row += 3
        ws.column_dimensions["A"].width = 120
    for ws in wb.worksheets[1:]:
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 60)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
