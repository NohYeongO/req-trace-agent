"""산출물 문서를 docx(Word)로 쓴다 — 공공 검수 제출 형식(표지 정보, 번호 붙은 절, 표, 확인란).

초안은 머리에 '초안(검토 전)'을 크게 표시하고, 확정본은 확인란에 확정자·일시를 채운다.
HWP는 지원하지 않는다(한계) — 한글 프로그램에서 docx를 열어 변환할 수 있다.
"""
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = "맑은 고딕"


def _font(run, size: float | None = None, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _para(doc, text: str, size: float = 10, bold: bool = False, align=None, color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    _font(p.add_run(text), size, bold, color)
    return p


def _table(doc, columns: list[str], rows: list[list]) -> None:
    t = doc.add_table(rows=1, cols=len(columns))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(columns):
        cell = t.rows[0].cells[i]
        cell.text = ""
        _font(cell.paragraphs[0].add_run(str(c)), 9, bold=True)
        _shade(cell, "DCE6F1")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            _font(cells[i].paragraphs[0].add_run("" if v is None else str(v)), 9)
    if not rows:
        cells = t.add_row().cells
        merged = cells[0].merge(cells[-1])
        merged.text = ""
        _font(merged.paragraphs[0].add_run("해당 없음"), 9)


def render_docx(record: dict) -> bytes:
    """record: 업무 DB의 deliverables 행(content·status·confirmed_by 등)."""
    content, confirmed = record["content"], record["status"] == "confirmed"
    meta = content["meta"]
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(10)

    if not confirmed:
        _para(doc, "초안(검토 전) — 담당자 확인 후 확정본을 제출하십시오", 11, True, WD_ALIGN_PARAGRAPH.CENTER,
              RGBColor(0xC0, 0x39, 0x2B))
    _para(doc, meta.get("project_name", ""), 12, False, WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, content["title"], 20, True, WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, f"수행 차수: {meta.get('run_label', '')}   작성일: {meta.get('written_on', '')}", 10, False,
          WD_ALIGN_PARAGRAPH.CENTER)

    for sec in content["sections"]:
        h = doc.add_paragraph()
        _font(h.add_run(sec["title"]), 13, True)
        if sec["kind"] == "table":
            _table(doc, sec["columns"], sec["rows"])
        else:
            for line in (sec.get("text") or "").splitlines() or [""]:
                _para(doc, line, 10)
            if sec["kind"] == "ai":
                origin = "담당자 수정" if sec.get("edited") else ("코드 작성(AI 서술 검증 실패)" if sec.get("fallback")
                                                              else "AI 초안 — 수치는 결과 자료와 대조함")
                _para(doc, f"※ {origin}", 8, color=RGBColor(0x80, 0x80, 0x80))
        if sec.get("note"):
            _para(doc, f"※ {sec['note']}", 8, color=RGBColor(0x80, 0x80, 0x80))

    h = doc.add_paragraph()
    _font(h.add_run("확인"), 13, True)
    _table(doc, ["구분", "성명", "일시", "서명"], [
        ["작성(시스템)", "req-trace-agent", record.get("created_at", "")[:16].replace("T", " "), ""],
        ["확정", record.get("confirmed_by") or "", (record.get("confirmed_at") or "")[:16].replace("T", " "), ""],
        ["검수", "", "", ""],
    ])
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
