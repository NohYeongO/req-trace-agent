"""요구사항 명세서 로더 — 엑셀(xlsx) 또는 JSON을 읽어 Requirement 목록으로 만든다.

기관마다 양식이 조금씩 달라서, 엑셀은 첫 행(헤더)에서 '요구사항ID/고유번호', '요구사항명',
'상세설명/세부내용' 같은 흔한 열 이름 변형을 받아들인다. 형식이 맞지 않는 행은 버리지 않고
경고로 돌려줘 업로드 화면에서 무엇이 빠졌는지 보여준다.
"""
import io
import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path

# 한글 조사가 바로 붙는 경우("SFR-001이랑")도 잡도록 \b 대신 영숫자 경계를 직접 쓴다.
REQ_ID_RE = re.compile(r"(?<![A-Za-z0-9])([A-Z]{3}-\d{3})(?!\d)")

_HEADER_ALIASES = {
    "id": ("요구사항id", "요구사항고유번호", "고유번호", "요구사항번호", "id"),
    "category": ("분류", "요구사항분류", "구분", "유형"),
    "name": ("요구사항명", "명칭", "요구사항명칭"),
    "description": ("상세설명", "세부내용", "요구사항상세설명", "정의", "설명"),
    "priority": ("우선순위", "중요도"),
}


@dataclass(frozen=True)
class Requirement:
    id: str
    category: str
    name: str
    description: str = ""
    priority: str = ""

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass
class LoadResult:
    requirements: list[Requirement] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _norm(h) -> str:
    return re.sub(r"\s+", "", str(h or "")).lower()


def _cell(v) -> str:
    return "" if v is None else str(v).strip()


def _collect(rows: list[dict], where: str, first_row: int) -> LoadResult:
    """first_row: 경고에 표시할 첫 데이터 행 번호(엑셀은 헤더 다음인 2행)."""
    result = LoadResult()
    seen: set[str] = set()
    for i, r in enumerate(rows, start=first_row):
        rid, name = r.get("id", ""), r.get("name", "")
        if not rid and not name:
            continue   # 빈 행
        if not REQ_ID_RE.fullmatch(rid):
            result.warnings.append(f"{where} {i}행: 고유번호 '{rid}' 형식이 아니어서 건너뜀 (예: SFR-001)")
            continue
        if not name:
            result.warnings.append(f"{where} {i}행: {rid} 요구사항명이 비어 있어 건너뜀")
            continue
        if rid in seen:
            result.warnings.append(f"{where} {i}행: {rid}가 중복되어 뒤의 행은 건너뜀")
            continue
        seen.add(rid)
        result.requirements.append(Requirement(rid, r.get("category", ""), name,
                                               r.get("description", ""), r.get("priority", "")))
    return result


def load_xlsx(data: bytes) -> LoadResult:
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("엑셀에 내용이 없습니다.")
    header = [_norm(h) for h in rows[0]]
    col: dict[str, int] = {}
    for key, aliases in _HEADER_ALIASES.items():
        idx = next((i for i, h in enumerate(header) if h in aliases), None)
        if idx is not None:
            col[key] = idx
    if "id" not in col or "name" not in col:
        raise ValueError("첫 행에 '요구사항ID(고유번호)'와 '요구사항명' 열이 있어야 합니다. "
                         f"읽은 헤더: {[h for h in rows[0] if h]}")
    parsed = [{k: _cell(row[i]) if i < len(row) else "" for k, i in col.items()}
              for row in rows[1:]]
    return _collect(parsed, "엑셀", first_row=2)


def load_json(data: bytes) -> LoadResult:
    payload = json.loads(data.decode("utf-8"))
    items = payload["requirements"] if isinstance(payload, dict) else payload
    rows = [{k: _cell(r.get(k)) for k in _HEADER_ALIASES} for r in items]
    return _collect(rows, "JSON", first_row=1)


def load_requirements(data: bytes, filename: str) -> LoadResult:
    suffix = Path(filename).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return load_xlsx(data)
    if suffix == ".json":
        return load_json(data)
    raise ValueError(f"지원하지 않는 형식입니다: {suffix or filename} (xlsx 또는 json)")


def load_requirements_file(path: Path) -> LoadResult:
    return load_requirements(path.read_bytes(), path.name)
