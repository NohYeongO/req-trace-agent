"""데모 요구사항 명세서 엑셀 생성 — data/requirements/*.json → 같은 이름의 .xlsx.

화면의 '명세서 업로드'에 그대로 올릴 수 있는 엑셀을 만든다(공공 SI 요구사항 정의서와 비슷한
열 구성). 실행: .venv/bin/python3 tools/build_demo_files.py
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
HEADERS = [("id", "요구사항ID", 12), ("category", "분류", 8), ("name", "요구사항명", 26),
           ("description", "상세설명", 70), ("priority", "우선순위", 9)]


def build_requirements_xlsx(src: Path, dst: Path) -> None:
    data = json.loads(src.read_text(encoding="utf-8"))
    wb = Workbook()
    ws = wb.active
    ws.title = "요구사항정의서"
    for i, (_, header, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=i, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        ws.column_dimensions[cell.column_letter].width = width
    for r, req in enumerate(data["requirements"], start=2):
        for c, (key, _, _) in enumerate(HEADERS, start=1):
            ws.cell(row=r, column=c, value=req.get(key, "")).alignment = Alignment(
                wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    wb.save(dst)


if __name__ == "__main__":
    for src in sorted((ROOT / "data" / "requirements").glob("*.json")):
        build_requirements_xlsx(src, src.with_suffix(".xlsx"))
        print(f"생성: {src.with_suffix('.xlsx').relative_to(ROOT)}")
