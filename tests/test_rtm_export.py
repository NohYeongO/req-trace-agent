"""검수용 추적표 엑셀 — 시트 구성, 대시보드와 같은 숫자, 기준일, 사람이 확인한 데이터만."""
from datetime import date

from openpyxl import load_workbook

from agent.dashboard import build_dashboard
from agent.rtm_export import TRACE_COLUMNS, export_rtm
from tools.seed_demo import seed

TODAY = date(2026, 3, 10)


def _cells(ws, col):
    return [ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)]


def test_workbook_matches_dashboard(store, tmp_path):
    seed(store)
    out = tmp_path / "rtm.xlsx"
    result = export_rtm(store, out, "데모 사업", today=TODAY)
    wb = load_workbook(out)
    assert wb.sheetnames == ["표지·요약", "요구사항 추적표", "완료 보고 이력"]

    summary = {r[0]: r[1] for r in wb["표지·요약"].iter_rows(min_row=3, max_row=9, values_only=True)}
    assert summary["사업명"] == "데모 사업" and summary["요구사항 수"] == 32
    assert summary["진척률(요구사항 기준)"].startswith("68.8%")
    assert summary["진척률(완료 조건 기준)"].startswith("77.8%")

    trace = wb["요구사항 추적표"]
    assert [c.value for c in trace[1]] == [h for h, _ in TRACE_COLUMNS]
    statuses = _cells(trace, 8)
    dash = build_dashboard(store, today=TODAY)
    assert {s: statuses.count(s) for s in dash["overall"]["counts"]} == dash["overall"]["counts"]
    assert result["overall"] == dash["overall"]
    assert trace.freeze_panes == "C2" and trace.auto_filter.ref == "A1:O33"

    row = {trace.cell(row=r, column=1).value: [c.value for c in trace[r]] for r in range(2, trace.max_row + 1)}
    pay = row["SFR-018"]
    assert pay[5].splitlines() == ["✔ 쿠폰 할인을 적용해 결제 금액을 계산한다", "✔ 결제 승인 후 주문을 확정한다",
                                   "✔ 결제에 실패하면 사용한 쿠폰을 복구한다", "☐ 실제 결제 대행사(PG)와 연동해 승인을 받는다"]
    assert (pay[6], pay[7], pay[9]) == ("3/4", "부분완료", "박지훈")
    assert "– [API] POST /api/orders/payment" in pay[10]   # 산출물 경로 없이 시드 → 확인 불가 표시
    assert row["SER-003"][9] is None and row["SER-003"][7] == "미착수"   # 보고 없음
    assert wb["완료 보고 이력"].max_row - 1 == 11


def test_as_of_and_unconfirmed_criteria(store, tmp_path):
    seed(store)
    store.replace_criteria("SER-003", [{"text": "AI 초안 조건", "origin": "ai"}], confirm=False)
    out = tmp_path / "rtm.xlsx"
    export_rtm(store, out, "데모", as_of="2026-02-20", today=TODAY)
    wb = load_workbook(out)
    trace = wb["요구사항 추적표"]
    row = {trace.cell(row=r, column=1).value: [c.value for c in trace[r]] for r in range(2, trace.max_row + 1)}
    assert row["SER-003"][5] == "(완료 조건 미확정)" and row["SER-003"][7] == "조건 미확정"
    assert row["SFR-019"][7] == "미착수"          # 3/7 보고는 기준일 이후라 반영 안 됨
    assert wb["완료 보고 이력"].max_row - 1 == 7
    assert wb["표지·요약"]["B4"].value == "2026-02-20"
