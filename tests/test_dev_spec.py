"""개발 명세서 — 소스 분석(엔티티·DTO·컨트롤러·DDL), 요구사항 연결, 이름 추정 검증, 작성 그래프, xlsx."""
import io
from datetime import date

from openpyxl import load_workbook

from agent.artifacts import ArtifactChecker
from agent.dashboard import build_dashboard
from agent.deliverables import dev_spec
from agent.deliverables.java_source import parse_ddl, parse_dto, parse_entity, snake
from agent.deliverables.xlsx_writer import render_xlsx
from agent.graphs import build_deliverable_graph
from agent.state import initial_deliverable_state
from tests.conftest import FIXTURES
from tools.seed_demo import seed

MINI = FIXTURES / "mini_spring"


def test_parse_entity_and_inheritance():
    spec = ArtifactChecker(MINI).source_spec()
    ents = {e["table"]: e for e in spec["entities"]}
    assert set(ents) == {"carts", "cart_items"}                     # BaseEntity(@MappedSuperclass)는 표에서 빠짐
    cols = {c["name"]: c for c in ents["cart_items"]["columns"]}
    assert list(cols)[:2] == ["id", "created_at"]                   # 상속 컬럼이 앞에
    assert cols["id"]["pk"] and cols["id"]["generated"] == "GenerationType.IDENTITY"
    assert cols["cart_id"]["fk"] == "Cart" and not cols["cart_id"]["nullable"]   # @JoinColumn(nullable=false)
    assert not cols["quantity"]["nullable"]                         # 기본형 int → NOT NULL
    assert cols["memo"]["length"] == 40 and cols["status"]["note"] == "enum ItemStatus"
    assert "MAX" not in cols and "local" not in cols                # static·메서드 지역변수 제외
    assert ents["carts"]["relations"] == ["OneToMany items: List<CartItem>"]
    assert ents["cart_items"]["unique_constraints"] == ["uk_cart_product(cart_id, product_id)"]
    assert ents["cart_items"]["issues"] == ["quantity: 엔티티는 NOT NULL, DDL은 NULL 허용",
                                            "memo: DDL에 없음(엔티티 필드 memo)",
                                            "legacy_flag: 엔티티에 없음(DDL에만 있음)"]
    assert spec["ddl_only"] == ["audit_log"] and spec["ddl_files"] == ["db/schema.sql"]
    assert snake("postalCode") == "postal_code" and parse_entity("class X {}") is None


def test_parse_controller_dto_ddl():
    spec = ArtifactChecker(MINI).source_spec()
    add, upd = spec["endpoints"]
    assert (add["method"], add["path"], add["summary"], add["auth"], add["body"]) == \
        ("POST", "/api/carts/items", "장바구니 상품 추가", True, "CartItemAddRequest")
    assert add["tag"] == "Cart" and add["returns"] == "ApiResult<CartResponse>"
    assert upd["summary"] == "" and upd["params"] == [
        {"in": "path", "name": "cartItemId", "type": "Long", "required": True, "description": "장바구니 항목 ID"},
        {"in": "query", "name": "dryRun", "type": "boolean", "required": False, "description": ""}]
    req = spec["dtos"]["CartItemAddRequest"]
    assert req["description"] == "장바구니 상품 추가 요청"
    assert req["fields"][1] == {"name": "quantity", "type": "Integer", "required": True,
                                "constraints": "NotNull, Min(1), Max(99)", "description": "수량", "example": ""}
    assert parse_dto("public record R(Long id, List<String> tags) {}")["fields"][1]["type"] == "List<String>"
    ddl = parse_ddl((MINI / "db" / "schema.sql").read_text())
    assert ddl["carts"]["columns"]["account_id"] == {"type": "BIGINT", "nullable": False, "default": "", "pk": False,
                                                     "unique": False, "auto": False, "comment": "계정 ID"}
    assert ddl["cart_items"]["columns"]["cart_id"]["fk"] == "carts.id"


def _doc(store):
    seed(store)
    prog = build_dashboard(store, today=date(2026, 3, 10))
    prog.pop("trend")
    spec = ArtifactChecker(MINI).source_spec()
    return dev_spec.build(spec, prog, store.list_reports(1000), {"project_name": "데모", "run_label": "abc1234"})


def test_links_tasks_and_name_checks(store):
    doc = _doc(store)
    eps = doc["model"]["spec"]["endpoints"]
    # 완료 보고 API 근거(POST /api/carts/items, PATCH /api/carts/items/{cartItemId}) → SFR-009·SFR-010
    assert [ep["if_id"] for ep in eps] == ["IF-CART-001", "IF-CART-002"]
    assert eps[0]["reqs"] == ["SFR-009", "SFR-010"]
    assert doc["tasks"] == ["names:cart_items", "names:carts", "endpoints", "opinion"]
    items, keys = dev_spec._table_items(doc, "carts")
    assert keys == ["carts", "carts.id", "carts.created_at"]       # DDL COMMENT가 있는 account_id는 묻지 않음
    assert dev_spec.check_task(doc, "names:carts", {"carts": "장바구니", "carts.id": "ID", "carts.created_at": "생성 일시"}) == []
    bad = dev_spec.check_task(doc, "names:carts", {"carts": "cart", "carts.id": "ID"})
    assert bad[0] == "이름이 빠진 key: carts.created_at" and "carts: 'cart'" in bad[1]
    assert dev_spec.check_task(doc, "opinion", "인터페이스 99개이다.") == ["사실에 없는 숫자 99"]


async def test_graph_builds_dev_spec(deps):
    deps.mcp.checker = ArtifactChecker(MINI)
    seed(deps.store)
    r = await build_deliverable_graph(deps).ainvoke(initial_deliverable_state("repo", "dev_spec"))
    assert r["doc_id"] and not r["fallbacks"], r.get("error")
    doc = deps.store.get_deliverable(r["doc_id"])["content"]
    sec = {s["id"]: s for s in doc["sections"]}
    assert sec["interfaces"]["rows"][0][:4] == ["IF-CART-001", "POST", "/api/carts/items", "장바구니 상품 추가"]
    assert sec["interfaces"]["rows"][1][3] == "API 기능 (추정)"           # 설명이 없어 AI가 추정
    assert sec["if_1"]["rows"][0] == ["요청 본문", "productId", "Long", "Y", "NotNull", "상품 ID"]
    assert sec["if_2"]["rows"][1] == ["쿼리", "dryRun", "boolean", "N", "", ""]
    tables = {r[0]: r for r in sec["tables"]["rows"]}
    assert tables["cart_items"][5] == "불일치 3건" and tables["carts"][5] == "일치"
    carts = {r[1]: r for r in sec["tb_2"]["rows"]}
    assert carts["account_id"][2] == "계정 ID"                           # DDL COMMENT는 그대로(추정 아님)
    assert carts["id"][2] == "ID (추정)" and carts["id"][6] == "PK"
    items = {r[1]: r for r in sec["tb_1"]["rows"]}
    assert items["cart_id"][7] == "carts.id" and items["memo"][3] == "VARCHAR" and "DDL에 없음" in items["memo"][9]
    assert ["", "엔티티 없음(DDL에만 있는 테이블)"] != sec["mismatch"]["rows"][-1]
    assert sec["mismatch"]["rows"][-1] == ["audit_log", "엔티티 없음(DDL에만 있는 테이블)"]
    assert sec["opinion"]["kind"] == "ai" and "인터페이스 2개" in sec["opinion"]["text"]
    assert "model" not in doc


async def test_dev_spec_without_repo(deps):
    deps.mcp.checker = ArtifactChecker(None)
    r = await build_deliverable_graph(deps).ainvoke(initial_deliverable_state("repo", "dev_spec"))
    assert "설정되지 않았습니다" in r["error"]


async def test_xlsx_render(deps):
    deps.mcp.checker = ArtifactChecker(MINI)
    seed(deps.store)
    r = await build_deliverable_graph(deps).ainvoke(initial_deliverable_state("repo", "dev_spec"))
    wb = load_workbook(io.BytesIO(render_xlsx(deps.store.get_deliverable(r["doc_id"]))))
    assert wb.sheetnames == ["표지", "개요", "인터페이스 목록", "인터페이스 상세", "테이블 목록", "테이블 정의",
                             "엔티티와 DDL 불일치", "구현 근거와 연결되지 않은 요구사항", "종합"]
    detail = [c.value for c in wb["인터페이스 상세"]["A"] if c.value]
    assert detail[0].startswith("2-1. IF-CART-001") and any(str(v).startswith("2-2. IF-CART-002") for v in detail)
    assert wb["표지"]["A3"].value.startswith("상태: 초안")
