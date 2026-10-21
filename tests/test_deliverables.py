"""검수 산출물 — 결과 파일 해석, 보고서 표 계산, 요구사항 연결, 작성 그래프, docx, API."""
import io
import json
import zipfile
from datetime import date

import pytest
from docx import Document

from agent.artifacts import ArtifactChecker
from agent.dashboard import build_dashboard
from agent.deliverables import sources, test_report
from agent.deliverables.docx_writer import render_docx
from agent.deliverables.junit import read_results, summarize, total_coverage
from agent.graphs import build_deliverable_graph
from agent.state import initial_deliverable_state
from tools.seed_demo import seed

RUN1 = sources.DEMO_DIR / "1차-2026-03-09"     # Docker 호환 문제로 통합 테스트 8건 실패
RUN2 = sources.DEMO_DIR / "2차-2026-03-10"     # 조치 후 전부 통과 + JaCoCo


def test_read_demo_results():
    r1, r2 = read_results(RUN1), read_results(RUN2)
    s1, s2 = summarize(r1["cases"]), summarize(r2["cases"])
    assert (s1["total"], s1["통과"], s1["실패"], s1["pass_rate"]) == (34, 26, 8, 76.5)
    assert (s2["total"], s2["통과"], s2["pass_rate"]) == (34, 34, 100.0)
    assert r1["coverage"] == [] and len(r2["coverage"]) == 5
    assert total_coverage(r2["coverage"])["line"] == {"covered": 200, "missed": 558, "rate": 26.4}
    fail = next(c for c in r1["cases"] if c["status"] == "실패")
    assert fail["class"] == "AddressPortServiceTest" and fail["module"] == "module-account"
    assert "Could not find a valid Docker environment" in fail["causes"][-1]
    nested = {c["class"] for c in r2["cases"] if "$" in c["classname"]}
    assert "CartAppServiceTest" in nested                     # 중첩 클래스(@Nested)는 바깥 클래스로 묶는다


def test_test_kinds_from_source(tmp_path):
    t = tmp_path / "src" / "test" / "java"
    t.mkdir(parents=True)
    (t / "OrderRepoTest.java").write_text("@SpringBootTest\nclass OrderRepoTest extends MysqlRedisTestContainer {}")
    (t / "CartServiceTest.java").write_text("class CartServiceTest { @Test void add() {} }")
    kinds = ArtifactChecker(tmp_path).test_kinds(["OrderRepoTest", "CartServiceTest", "Nope"])
    assert kinds["OrderRepoTest"]["kind"] == "통합" and kinds["CartServiceTest"]["kind"] == "단위"
    assert kinds["Nope"]["kind"] == "미확인"


@pytest.fixture
def seeded(store):
    seed(store)
    prog = build_dashboard(store, today=date(2026, 3, 10))
    prog.pop("trend")
    return prog, store.list_reports(1000)


KINDS = {"AddressPortServiceTest": {"kind": "통합"}, "CouponPortServiceTest": {"kind": "통합"},
         "OrderSessionRepositoryTest": {"kind": "통합"}}


def test_build_report_tables(seeded):
    prog, reports = seeded
    res = read_results(RUN1)
    kinds = {c: KINDS.get(c, {"kind": "단위"}) for c in {x["class"] for x in res["cases"]}}
    doc = test_report.build(res, kinds, prog, reports, {"project_name": "데모", "run_label": "1차"})
    sec = {s["id"]: s for s in doc["sections"]}
    assert sec["summary"]["rows"] == [["단위", 26, 26, 0, 0, 0, "100.0%"], ["통합", 8, 0, 8, 0, 0, "0.0%"],
                                      ["합계", 34, 26, 8, 0, 0, "76.5%"]]
    assert sec["coverage"]["kind"] == "text" and "자료 없음" in sec["coverage"]["text"]
    assert len(sec["failures"]["rows"]) == 8 and sec["failure_analysis"]["kind"] == "ai"
    assert sec["failure_analysis"]["groups"][0]["count"] == 8           # 원인 하나로 묶임
    links = {r[0]: r for r in sec["req_links"]["rows"]}
    # 완료 보고의 테스트 근거: CartAppServiceTest → SFR-009·SFR-010
    assert links["SFR-009"][3] == "CartAppServiceTest" and links["SFR-009"][5] == "통과"
    assert links["SFR-009"][6].startswith("완료 보고 #")
    # 명명 규칙: 보고의 코드 근거 'CouponPortService.cancelUsage' → CouponPortServiceTest
    assert "CouponPortServiceTest" in links["SFR-018"][3] and "명명 규칙" in links["SFR-018"][6]
    # 보고 #8의 테스트 근거 2개(OrderSessionRepositoryTest 3건, CouponPortServiceTest 3건)가 모두 실패
    assert links["SFR-016"][3] == "OrderSessionRepositoryTest, CouponPortServiceTest" and links["SFR-016"][5] == "실패 6건"
    untested = [r[0] for r in sec["untested"]["rows"]]
    assert "SFR-019" in untested and "SFR-009" not in untested
    assert [r[0] for r in sec["orphans"]["rows"]] == ["AddressPortServiceTest"]
    assert "[실패 원인] java.lang.IllegalStateException: Could not find a valid Docker environment" in doc["facts"]


async def test_graph_writes_validates_and_saves(deps):
    seed(deps.store)
    g = build_deliverable_graph(deps)
    r = await g.ainvoke(initial_deliverable_state("demo:1차-2026-03-09"))
    assert r["doc_id"] and not r["fallbacks"]
    saved = deps.store.get_deliverable(r["doc_id"])
    assert saved["status"] == "draft" and saved["title"] == "테스트 결과 보고서 (1차-2026-03-09)"
    sec = {s["id"]: s for s in saved["content"]["sections"]}
    assert "8건" in sec["failure_analysis"]["text"] and "76.5%" in sec["opinion"]["text"]
    assert set(deps.mcp.calls) >= {"analyze_test_results", "get_progress", "list_reports"}
    types = [e["type"] for e in r["events"]]
    assert types == ["ResultsCollected", "DocumentBuilt", "SectionsWritten", "SectionsValidated", "DraftSaved"]


async def test_invented_numbers_fall_back_to_code_text(deps):
    class Liar:
        name = "liar"

        async def write_section(self, kind, facts, feedback=None, on_token=None):
            return "전체 통과율은 99.9%이다.", {"usd": 0.0}

    seed(deps.store)
    deps.quality_llm = Liar()
    r = await build_deliverable_graph(deps).ainvoke(initial_deliverable_state("demo:2차-2026-03-10"))
    assert r["attempts"] == 3 and r["fallbacks"] == ["opinion"]     # 2차는 실패가 없어 의견 섹션만
    opinion = next(s for s in deps.store.get_deliverable(r["doc_id"])["content"]["sections"] if s["id"] == "opinion")
    assert "99.9" not in opinion["text"] and opinion["text"].startswith("전체 34건 중 34건이 통과했다(통과율 100.0%).")
    assert opinion["fallback"]


async def test_missing_source_stops_early(deps):
    r = await build_deliverable_graph(deps).ainvoke(initial_deliverable_state("demo:없는자료"))
    assert "찾을 수 없습니다" in r["error"] and not r.get("doc_id")


def test_docx_render(seeded, store):
    prog, reports = seeded
    res = read_results(RUN2)
    doc = test_report.build(res, {}, prog, reports, {"project_name": "데모 사업", "run_label": "2차"})
    for s in doc["sections"]:
        if s["kind"] == "ai":
            s["text"] = "검토용 문장이다."
    doc_id = store.save_deliverable(doc["doc_type"], "테스트 결과 보고서 (2차)", "demo:2차", doc)
    draft = Document(io.BytesIO(render_docx(store.get_deliverable(doc_id))))
    text = "\n".join(p.text for p in draft.paragraphs)
    assert text.startswith("초안(검토 전)") and "2. 수행 결과 요약" in text and "3. 코드 커버리지" in text
    assert draft.tables[1].rows[-1].cells[-1].text == "100.0%"          # 요약 표 합계 통과율
    store.confirm_deliverable(doc_id, "이서연")
    final = Document(io.BytesIO(render_docx(store.get_deliverable(doc_id))))
    assert not final.paragraphs[0].text.startswith("초안")
    assert final.tables[-1].rows[2].cells[1].text == "이서연"            # 확인란: 확정자
    with pytest.raises(ValueError):
        store.update_deliverable_text(doc_id, {"opinion": "수정"})       # 확정본은 수정 불가


def test_upload_zip_rejects_traversal(settings, tmp_path):
    from dataclasses import replace
    st = replace(settings, upload_dir=tmp_path / "up")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("module-a/TEST-x.xml", '<testsuite name="x" tests="1"><testcase name="t" classname="a.XTest"/></testsuite>')
        z.writestr("../../evil.xml", "<x/>")
        z.writestr("readme.txt", "skip")
    source = sources.save_upload(buf.getvalue(), "results.zip", st)
    path = sources.resolve(source, st)
    assert sorted(p.name for p in path.rglob("*.xml")) == ["TEST-x.xml", "evil.xml"]
    assert all(path in p.parents for p in path.rglob("*.xml"))           # 폴더 밖으로 못 나감
    assert not (tmp_path / "evil.xml").exists()
    with pytest.raises(ValueError):
        sources.save_upload(b"not zip", "x.zip", st)
    with pytest.raises(ValueError):
        sources.resolve("upload:../../etc", st)
